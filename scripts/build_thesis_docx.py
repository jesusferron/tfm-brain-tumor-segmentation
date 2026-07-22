#!/usr/bin/env python3
"""Build the final TFM manuscript on top of the official VIU DOCX template.

The Markdown files remain the source of truth. Pandoc is used for the body so
that tables, links and LaTeX equations become native Word content. The cover,
section layout, headers, footers and styles come from the supplied template.
"""

from __future__ import annotations

import argparse
import re
import tempfile
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import pypandoc
from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph
from docxcompose.composer import Composer
from lxml import etree


ROOT = Path(__file__).resolve().parents[1]
MEMORIA = ROOT / "docs" / "memoria"
DEFAULT_OUTPUT = MEMORIA / "TFM_memoria_maquetada_borrador.docx"

TITLE = (
    "Segmentación de tumores cerebrales en resonancia magnética multimodal "
    "mediante arquitecturas Transformer-UNet híbridas"
)

SOURCE_FILES = [
    MEMORIA / "resumen-abstract.md",
    MEMORIA / "capitulo-1-introduccion.md",
    MEMORIA / "capitulo-2-marco-teorico.md",
    MEMORIA / "capitulo-3-metodologia.md",
    MEMORIA / "capitulo-4-desarrollo.md",
    MEMORIA / "capitulo-5-resultados.md",
    MEMORIA / "capitulo-6-conclusiones.md",
    MEMORIA / "capitulo-7-referencias.md",
    MEMORIA / "apendice-reproducibilidad.md",
]

FIGURE_ALT_TEXTS = [
    "Flujo global del sistema experimental.",
    "Arquitectura de fusión multimodal FusionUNet y detalle de la compuerta adaptativa.",
    "Comparación cualitativa de arquitecturas: referencia, nnU-Net, Swin-UNETR, "
    "Attention U-Net y Residual U-Net con concatenación.",
    "Comparación entre la concatenación y una corrida de bajo rendimiento de "
    "la compuerta adaptativa.",
]

CAPTION_RE = re.compile(r"^(Tabla|Figura)\s+((?:\d+)|(?:A\.1))\.\s*(.+)$", re.DOTALL)
IMAGE_RE = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--template",
        type=Path,
        required=True,
        help="Ruta a la plantilla DOCX oficial de la universidad.",
    )
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--title", default=TITLE)
    parser.add_argument("--degree", default="PENDIENTE")
    parser.add_argument("--academic-year", default="PENDIENTE")
    parser.add_argument("--student", default="PENDIENTE")
    parser.add_argument("--dni", default="PENDIENTE")
    parser.add_argument("--director", default="PENDIENTE")
    parser.add_argument("--call", default="PENDIENTE")
    parser.add_argument("--date", default="PENDIENTE")
    return parser.parse_args()


def iter_block_items(parent: DocumentObject):
    """Yield top-level paragraphs and tables in document order."""
    parent_elm = parent.element.body
    for child in parent_elm.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def set_run_font(run, name: str, size: float | None = None) -> None:
    run.font.name = name
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
        fonts.set(qn(f"w:{attr}"), name)
    if size is not None:
        run.font.size = Pt(size)


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    """Replace text while retaining the formatting of the first cover run."""
    if paragraph.runs:
        first = paragraph.runs[0]
        first.text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def clear_paragraph_content(paragraph: Paragraph) -> None:
    p = paragraph._p
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)


def set_language(paragraph: Paragraph, language: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    ppr_rpr = ppr.find(qn("w:rPr"))
    if ppr_rpr is None:
        ppr_rpr = OxmlElement("w:rPr")
        ppr.append(ppr_rpr)
    lang = ppr_rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        ppr_rpr.append(lang)
    lang.set(qn("w:val"), language)
    for run in paragraph.runs:
        rpr = run._r.get_or_add_rPr()
        run_lang = rpr.find(qn("w:lang"))
        if run_lang is None:
            run_lang = OxmlElement("w:lang")
            rpr.append(run_lang)
        run_lang.set(qn("w:val"), language)


def add_complex_field(
    paragraph: Paragraph,
    instruction: str,
    result: str,
    *,
    bold: bool = False,
    placeholder: bool = False,
) -> None:
    """Append an updateable Word field with a useful cached result."""
    begin_run = OxmlElement("w:r")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    begin_run.append(begin)

    instruction_run = OxmlElement("w:r")
    instruction_text = OxmlElement("w:instrText")
    instruction_text.set(qn("xml:space"), "preserve")
    instruction_text.text = f" {instruction} "
    instruction_run.append(instruction_text)

    separator_run = OxmlElement("w:r")
    separator = OxmlElement("w:fldChar")
    separator.set(qn("w:fldCharType"), "separate")
    separator_run.append(separator)

    result_run = OxmlElement("w:r")
    result_rpr = OxmlElement("w:rPr")
    if bold:
        result_rpr.append(OxmlElement("w:b"))
    if placeholder:
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "7F7F7F")
        result_rpr.append(color)
        result_rpr.append(OxmlElement("w:i"))
    result_run.append(result_rpr)
    result_text = OxmlElement("w:t")
    result_text.text = result
    result_run.append(result_text)

    end_run = OxmlElement("w:r")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run.append(end)

    for element in (begin_run, instruction_run, separator_run, result_run, end_run):
        paragraph._p.append(element)


def set_no_number(paragraph: Paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    old = ppr.find(qn("w:numPr"))
    if old is not None:
        ppr.remove(old)
    num_pr = OxmlElement("w:numPr")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), "0")
    num_pr.append(num_id)
    ppr.insert(0, num_pr)


def remove_numbering(paragraph: Paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = ppr.find(qn("w:numPr"))
    if num_pr is not None:
        ppr.remove(num_pr)


def shade(element, fill: str) -> None:
    props = element.get_or_add_tcPr() if hasattr(element, "get_or_add_tcPr") else element
    old = props.find(qn("w:shd"))
    if old is not None:
        props.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    props.append(shd)


def add_paragraph_shading(paragraph: Paragraph, fill: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    old = ppr.find(qn("w:shd"))
    if old is not None:
        ppr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def add_left_border(paragraph: Paragraph, color: str, size: int = 18) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    borders = ppr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        ppr.append(borders)
    left = borders.find(qn("w:left"))
    if left is None:
        left = OxmlElement("w:left")
        borders.append(left)
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), color)


def prepare_markdown() -> str:
    missing = [path for path in SOURCE_FILES if not path.exists()]
    if missing:
        raise FileNotFoundError(f"Missing manuscript sources: {missing}")
    chunks = []
    for path in SOURCE_FILES:
        text = path.read_text(encoding="utf-8").strip()
        if path.name == "capitulo-7-referencias.md":
            # Collapse hard-wrapped bibliography items before Pandoc parses
            # them. A continuation such as "  1000587. https://..." would
            # otherwise be misread as a nested ordered-list item.
            lines = text.splitlines()
            normalized: list[str] = []
            current: str | None = None
            for line in lines:
                if line.startswith("- "):
                    if current is not None:
                        normalized.append(current)
                    current = line
                elif current is not None and (line.startswith("  ") or not line.strip()):
                    if line.strip():
                        current += " " + line.strip()
                else:
                    if current is not None:
                        normalized.append(current)
                        current = None
                    normalized.append(line)
            if current is not None:
                normalized.append(current)
            text = "\n".join(normalized)
        # A standalone image with non-empty alt text becomes a Pandoc figure and
        # would duplicate the explicit caption already present in the sources.
        text = IMAGE_RE.sub(lambda match: f"![]({match.group(2)})", text)
        chunks.append(text)
    return "\n\n".join(chunks) + "\n"


def pandoc_body(markdown: str, template: Path, output: Path) -> None:
    pypandoc.convert_text(
        markdown,
        to="docx",
        format=(
            "markdown+tex_math_dollars+pipe_tables+fenced_code_blocks+"
            "autolink_bare_uris"
        ),
        outputfile=str(output),
        extra_args=[
            f"--reference-doc={template}",
            f"--resource-path={MEMORIA}",
            "--standalone",
            "--wrap=none",
            "--metadata=lang:es-ES",
        ],
    )


def trim_template_and_fill_cover(template: Path, output: Path, args: argparse.Namespace) -> None:
    document = Document(template)
    body = document._element.body
    final_section = body.sectPr
    cover_end = None
    for child in body.iterchildren():
        if child.tag == qn("w:p") and child.find("./w:pPr/w:sectPr", child.nsmap) is not None:
            cover_end = child
            break
    if cover_end is None:
        raise RuntimeError("Could not find the cover section break in the template")

    after_cover = False
    for child in list(body):
        if child is cover_end:
            after_cover = True
            continue
        if after_cover and child is not final_section:
            body.remove(child)

    title_paragraph = next(p for p in document.paragraphs if p.text == "Título del TFT")
    set_paragraph_text(title_paragraph, args.title)
    title_run = title_paragraph.runs[0]
    title_run.bold = True
    title_run.font.highlight_color = None
    set_run_font(title_run, "Arial", 24)
    title_paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    cover = document.tables[0]
    cells = cover.rows[0].cells
    set_paragraph_text(cells[0].paragraphs[1], args.degree)
    set_paragraph_text(cells[0].paragraphs[3], args.academic_year)
    set_paragraph_text(cells[1].paragraphs[1], args.student)
    set_paragraph_text(cells[1].paragraphs[2], f"D.N.I.: {args.dni}")
    set_paragraph_text(cells[1].paragraphs[4], f"Director/a de TFM: {args.director}")
    set_paragraph_text(cells[2].paragraphs[2], args.call)
    # The original middle cell contains an empty spacer paragraph. With the
    # longer placeholder value it pushes the director line below the floating
    # cover table in LibreOffice (and can do the same in Word). Remove only
    # that spacer; all labels and the table geometry remain untouched.
    middle_spacer = cells[1].paragraphs[3]
    middle_spacer._element.getparent().remove(middle_spacer._element)
    for cell in cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                set_run_font(run, "Arial", 9.5)
                run.font.color.rgb = RGBColor(255, 255, 255)

    date_paragraph = next(p for p in document.paragraphs if p.text == "00 Mes 2023")
    set_paragraph_text(date_paragraph, args.date)
    for run in date_paragraph.runs:
        set_run_font(run, "Arial", 10)

    document.save(output)


def ensure_style(document: DocumentObject, name: str, base: str):
    try:
        return document.styles[name]
    except KeyError:
        style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = document.styles[base]
        return style


def configure_styles(document: DocumentObject) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_after = Pt(10)
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size in (("Heading 1", 22), ("Heading 2", 16), ("Heading 3", 12)):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        num_pr = style.element.find("./w:pPr/w:numPr", style.element.nsmap)
        if num_pr is not None:
            num_pr.getparent().remove(num_pr)
    document.styles["Heading 3"].paragraph_format.space_before = Pt(10)
    document.styles["Heading 3"].paragraph_format.space_after = Pt(5)

    caption = document.styles["Caption"]
    caption.font.name = "Arial"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor(0, 0, 0)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.line_spacing = 1.0

    code = ensure_style(document, "TFM Code", "Normal")
    code.font.name = "Consolas"
    code.font.size = Pt(7.5)
    code.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    code.paragraph_format.left_indent = Cm(0.25)
    code.paragraph_format.right_indent = Cm(0.15)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(6)
    code.paragraph_format.line_spacing = 1.0

    bibliography = ensure_style(document, "TFM Bibliography", "Normal")
    bibliography.font.name = "Arial"
    bibliography.font.size = Pt(10)
    bibliography.paragraph_format.left_indent = Cm(1.27)
    bibliography.paragraph_format.first_line_indent = Cm(-1.27)
    bibliography.paragraph_format.space_after = Pt(6)
    bibliography.paragraph_format.line_spacing = 1.0

    highlight = ensure_style(document, "TFM Highlight", "Normal")
    highlight.font.name = "Arial"
    highlight.font.size = Pt(11)
    highlight.font.italic = True
    highlight.paragraph_format.left_indent = Cm(0.7)
    highlight.paragraph_format.right_indent = Cm(0.7)
    highlight.paragraph_format.space_before = Pt(8)
    highlight.paragraph_format.space_after = Pt(8)
    highlight.paragraph_format.line_spacing = 1.15


def insert_indices(document: DocumentObject) -> None:
    target = next(
        p
        for p in document.paragraphs
        if p.style.name == "Heading 1" and p.text.strip() == "1. Introducción"
    )
    entries = [
        ("Índice general", 'TOC \\o "1-3" \\h \\z \\u'),
        ("Índice de figuras", 'TOC \\c "Figura" \\h \\z'),
        ("Índice de tablas", 'TOC \\c "Tabla" \\h \\z'),
    ]
    for title, instruction in entries:
        heading = target.insert_paragraph_before(title, style="TOC Heading")
        heading.paragraph_format.page_break_before = True
        heading.paragraph_format.keep_with_next = True
        heading.paragraph_format.space_after = Pt(12)
        field_paragraph = target.insert_paragraph_before()
        field_paragraph.style = document.styles["Normal"]
        field_paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_complex_field(
            field_paragraph,
            instruction,
            "Actualizar este índice en Microsoft Word (Ctrl+A y F9).",
            placeholder=True,
        )


def move_table_captions_below_tables(document: DocumentObject) -> None:
    for paragraph in list(document.paragraphs):
        if not paragraph.text.startswith("Tabla "):
            continue
        if CAPTION_RE.match(paragraph.text.strip()) is None:
            continue
        sibling = paragraph._p.getnext()
        while sibling is not None and sibling.tag == qn("w:p"):
            sibling_paragraph = Paragraph(sibling, paragraph._parent)
            if sibling_paragraph.text.strip():
                break
            sibling = sibling.getnext()
        if sibling is not None and sibling.tag == qn("w:tbl"):
            sibling.addnext(paragraph._p)


def build_caption(paragraph: Paragraph) -> None:
    match = CAPTION_RE.match(paragraph.text.strip())
    if match is None:
        return
    label, number, description = match.groups()
    clear_paragraph_content(paragraph)
    paragraph.style = paragraph.part.document.styles["Caption"]
    paragraph.paragraph_format.keep_together = True

    prefix = paragraph.add_run(f"{label} ")
    prefix.bold = True
    if number == "A.1":
        appendix_prefix = paragraph.add_run("A.")
        appendix_prefix.bold = True
        add_complex_field(
            paragraph,
            "SEQ Tabla \\r 1 \\* ARABIC",
            "1",
            bold=True,
        )
    else:
        add_complex_field(
            paragraph,
            f"SEQ {label} \\* ARABIC",
            number,
            bold=True,
        )
    punctuation = paragraph.add_run(". ")
    punctuation.bold = True
    paragraph.add_run(description)


def configure_headings_and_paragraphs(document: DocumentObject) -> None:
    in_manuscript = False
    in_abstract = False
    in_references = False
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        style_name = paragraph.style.name

        if style_name == "Heading 1" and text == "Resumen":
            in_manuscript = True
        if not in_manuscript:
            continue

        if style_name == "Heading 1":
            in_abstract = text == "Abstract"
            in_references = text == "Referencias bibliográficas"
            if text.startswith("Apéndice A"):
                in_references = False
            set_no_number(paragraph)
            paragraph.paragraph_format.page_break_before = text != "Resumen"
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.keep_together = True
        elif style_name in ("Heading 2", "Heading 3"):
            set_no_number(paragraph)
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.keep_together = True

        if style_name == "TOC Heading":
            in_abstract = False
            in_references = False
            continue

        if CAPTION_RE.match(text):
            build_caption(paragraph)
        elif style_name == "Source Code":
            paragraph.style = document.styles["TFM Code"]
            paragraph.paragraph_format.keep_together = False
            paragraph.paragraph_format.widow_control = False
            add_paragraph_shading(paragraph, "F2F2F2")
            add_left_border(paragraph, "ED7D31", 12)
            for run in paragraph.runs:
                set_run_font(run, "Consolas", 7.5)
                run.font.color.rgb = RGBColor(31, 31, 31)
                rpr = run._r.get_or_add_rPr()
                if rpr.find(qn("w:noProof")) is None:
                    rpr.append(OxmlElement("w:noProof"))
        elif style_name == "Block Text" or text.startswith("¿Puede una ponderación explícita"):
            paragraph.style = document.styles["TFM Highlight"]
            add_paragraph_shading(paragraph, "FCE4D6")
            add_left_border(paragraph, "ED7D31")
        elif in_references and text and style_name not in ("Heading 1", "Heading 2", "Heading 3"):
            remove_numbering(paragraph)
            paragraph.style = document.styles["TFM Bibliography"]
        elif style_name not in ("Heading 1", "Heading 2", "Heading 3", "TOC Heading"):
            paragraph.style = document.styles["Normal"]

        if in_abstract:
            set_language(paragraph, "en-US")
        else:
            set_language(paragraph, "es-ES")

        # Pandoc character styles preserve semantics; force a compact
        # monospaced face only for inline code and syntax-highlighted tokens.
        for run in paragraph.runs:
            run_style = run.style.name if run.style is not None else ""
            if "Verbatim" in run_style or run_style.endswith("Tok"):
                set_run_font(run, "Consolas", 9 if style_name != "Source Code" else 7.5)
                rpr = run._r.get_or_add_rPr()
                if rpr.find(qn("w:noProof")) is None:
                    rpr.append(OxmlElement("w:noProof"))


def set_cell_margins(cell, value: int = 70) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side in ("top", "left", "bottom", "right"):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_table_borders(table: Table, color: str = "BFBFBF") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "5")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def format_tables(document: DocumentObject) -> None:
    content_tables = document.tables[1:]
    for table in content_tables:
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        tbl_pr = table._tbl.tblPr
        layout = tbl_pr.find(qn("w:tblLayout"))
        if layout is None:
            layout = OxmlElement("w:tblLayout")
            tbl_pr.append(layout)
        layout.set(qn("w:type"), "fixed")
        width = tbl_pr.find(qn("w:tblW"))
        if width is None:
            width = OxmlElement("w:tblW")
            tbl_pr.append(width)
        width.set(qn("w:w"), "8504")
        width.set(qn("w:type"), "dxa")
        set_table_borders(table)

        columns = max(len(row.cells) for row in table.rows)
        font_size = 7.2 if columns >= 6 else 8 if columns == 5 else 8.5
        column_width = Cm(15 / columns)
        set_repeat_header(table.rows[0])

        for row_index, row in enumerate(table.rows):
            prevent_row_split(row)
            for cell in row.cells:
                cell.width = column_width
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margins(cell)
                shade(cell._tc, "ED7D31" if row_index == 0 else ("FCE4D6" if row_index % 2 == 0 else "FFFFFF"))
                for paragraph in cell.paragraphs:
                    # Pandoc uses a ``Compact`` paragraph style for table
                    # cells. The official template does not define it, and
                    # some renderers then detach the text from the cell grid.
                    paragraph.style = document.styles["Normal"]
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.0
                    paragraph.paragraph_format.alignment = (
                        WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
                    )
                    for run in paragraph.runs:
                        set_run_font(run, "Arial", font_size)
                        run.font.bold = True if row_index == 0 else run.bold
                        run.font.color.rgb = (
                            RGBColor(255, 255, 255) if row_index == 0 else RGBColor(0, 0, 0)
                        )
            if row_index == len(table.rows) - 1:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.keep_with_next = True


def format_figures(document: DocumentObject) -> None:
    shapes = list(document.inline_shapes)
    if len(shapes) != len(FIGURE_ALT_TEXTS):
        raise RuntimeError(
            f"Expected {len(FIGURE_ALT_TEXTS)} manuscript figures, found {len(shapes)}"
        )
    for shape, alt_text in zip(shapes, FIGURE_ALT_TEXTS, strict=True):
        old_width = shape.width
        old_height = shape.height
        new_width = Cm(15)
        shape.width = new_width
        shape.height = int(old_height * int(new_width) / int(old_width))
        shape._inline.docPr.set("descr", alt_text)
        shape._inline.docPr.set("title", alt_text)
        ancestor = shape._inline
        while ancestor is not None and ancestor.tag != qn("w:p"):
            ancestor = ancestor.getparent()
        if ancestor is not None:
            paragraph = Paragraph(ancestor, document)
            paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(0)


def enable_field_updates(document: DocumentObject) -> None:
    settings = document.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def clean_properties(document: DocumentObject, title: str) -> None:
    props = document.core_properties
    props.title = title
    props.subject = "Trabajo Fin de Máster"
    props.author = ""
    props.last_modified_by = ""
    props.keywords = (
        "segmentación 3D, resonancia magnética, BraTS-GLI 2024, "
        "fusión multimodal, deep learning"
    )
    props.comments = ""


def clean_package_properties(path: Path, title: str) -> None:
    temp_path = path.with_suffix(".cleaning.docx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
        temp_path, "w", zipfile.ZIP_DEFLATED
    ) as target:
        for item in source.infolist():
            payload = source.read(item.filename)
            if item.filename == "docProps/core.xml":
                root = etree.fromstring(payload)
                core_ns = "http://schemas.openxmlformats.org/package/2006/metadata/core-properties"
                custom_ns = "http://schemas.openxmlformats.org/officeDocument/2006/custom-properties"
                dc_ns = "http://purl.org/dc/elements/1.1/"
                dcterms_ns = "http://purl.org/dc/terms/"
                xsi_ns = "http://www.w3.org/2001/XMLSchema-instance"

                # python-docx inherits a conflicting ``cp`` namespace mapping
                # from some older templates. Remove any malformed duplicates
                # and write the package metadata with explicit namespace URIs.
                for child in list(root):
                    if etree.QName(child).namespace == custom_ns:
                        root.remove(child)

                def set_core(namespace: str, local: str, value: str):
                    node = root.find(f"{{{namespace}}}{local}")
                    if node is None:
                        node = etree.SubElement(root, f"{{{namespace}}}{local}")
                    node.text = value
                    return node

                set_core(dc_ns, "title", title)
                set_core(dc_ns, "subject", "Trabajo Fin de Máster")
                set_core(dc_ns, "creator", "")
                set_core(dc_ns, "description", "")
                set_core(core_ns, "lastModifiedBy", "")
                set_core(core_ns, "revision", "1")
                set_core(
                    core_ns,
                    "keywords",
                    "segmentación 3D, resonancia magnética, BraTS-GLI 2024, "
                    "fusión multimodal, deep learning",
                )
                set_core(dc_ns, "language", "es-ES")
                timestamp = datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace(
                    "+00:00", "Z"
                )
                for local in ("created", "modified"):
                    date_node = set_core(dcterms_ns, local, timestamp)
                    date_node.set(f"{{{xsi_ns}}}type", "dcterms:W3CDTF")
                payload = etree.tostring(
                    root, xml_declaration=True, encoding="UTF-8", standalone=True
                )
            elif item.filename == "docProps/app.xml":
                root = etree.fromstring(payload)
                namespace = "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"
                company = root.find(f"{{{namespace}}}Company")
                if company is not None:
                    company.text = ""
                manager = root.find(f"{{{namespace}}}Manager")
                if manager is not None:
                    manager.text = ""
                payload = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            target.writestr(item, payload)
    temp_path.replace(path)


def validate(document: DocumentObject) -> dict[str, object]:
    body_tables = document.tables[1:]
    captions = [p.text.strip() for p in document.paragraphs if CAPTION_RE.match(p.text.strip())]
    headings = [p.text.strip() for p in document.paragraphs if p.style.name == "Heading 1"]
    xml = document._element
    instructions = " ".join(node.text or "" for node in xml.xpath(".//w:instrText"))
    extracted = "\n".join(p.text for p in document.paragraphs)
    style_counts: dict[str, int] = {}
    for paragraph in document.paragraphs:
        style_counts[paragraph.style.name] = style_counts.get(paragraph.style.name, 0) + 1
    external_relationships = [rel for rel in document.part.rels.values() if rel.is_external]
    expected_dimensions = [
        (4, 5), (7, 6), (6, 3), (6, 2), (6, 3), (4, 4), (4, 4), (6, 2),
        (8, 4), (8, 4), (10, 2), (5, 5), (8, 6), (8, 4), (8, 6), (9, 3),
    ]
    actual_dimensions = [(len(t.rows), max(len(r.cells) for r in t.rows)) for t in body_tables]
    checks = {
        "sections": len(document.sections),
        "tables": len(body_tables),
        "table_dimensions": actual_dimensions,
        "table_dimensions_ok": actual_dimensions == expected_dimensions,
        "figures": len(document.inline_shapes),
        "captions": len(captions),
        "seq_tabla": instructions.count("SEQ Tabla"),
        "seq_figura": instructions.count("SEQ Figura"),
        "toc_fields": instructions.count("TOC "),
        "math_total": len(xml.xpath(".//m:oMath")),
        "equations": len(xml.xpath(".//m:oMathPara")),
        "hyperlinks": len(xml.xpath(".//w:hyperlink")),
        "external_hyperlinks": len(external_relationships),
        "code_blocks": style_counts.get("TFM Code", 0),
        "bibliography_entries": style_counts.get("TFM Bibliography", 0),
        "heading_2": style_counts.get("Heading 2", 0),
        "heading_3": style_counts.get("Heading 3", 0),
        "headings_1": headings,
        "template_instructions_removed": "¡¡¡INFORMATIVO!!!" not in extracted,
        "latex_source_removed": all(
            marker not in extracted
            for marker in ("\\mathbf", "\\operatorname", "\\boldsymbol", "$$")
        ),
        "command_sentinels_ok": all(
            marker in extracted
            for marker in (
                "--fail-on-problems",
                "--ratios 70/15/15",
                "--max-steps 15000",
                '"$nnUNet_raw"',
                "checkpoint_best.pth",
            )
        ),
        "commit_marker_occurrences": extracted.count("PENDIENTE_COMMIT_FINAL"),
    }
    required = {
        "sections": 2,
        "tables": 16,
        "figures": 4,
        "captions": 20,
        "seq_tabla": 16,
        "seq_figura": 4,
        "toc_fields": 3,
        "math_total": 20,
        "equations": 4,
        "hyperlinks": 39,
        "external_hyperlinks": 39,
        "code_blocks": 10,
        "bibliography_entries": 34,
        "heading_2": 41,
        "heading_3": 25,
        "commit_marker_occurrences": 0,
    }
    failures = [key for key, expected in required.items() if checks[key] != expected]
    if not checks["table_dimensions_ok"]:
        failures.append("table_dimensions_ok")
    if not checks["template_instructions_removed"]:
        failures.append("template_instructions_removed")
    if not checks["latex_source_removed"]:
        failures.append("latex_source_removed")
    if not checks["command_sentinels_ok"]:
        failures.append("command_sentinels_ok")
    expected_headings = [
        "Resumen",
        "Abstract",
        "1. Introducción",
        "2. Marco teórico y estado del arte",
        "3. Metodología",
        "4. Desarrollo",
        "5. Resultados",
        "6. Conclusiones",
        "Referencias bibliográficas",
        "Apéndice A. Reproducibilidad",
    ]
    if headings != expected_headings:
        failures.append("headings_1")
    if failures:
        raise RuntimeError(f"DOCX validation failed: {failures}; details={checks}")
    return checks


def build(args: argparse.Namespace) -> dict[str, object]:
    template = args.template.expanduser().resolve()
    output = args.output.expanduser().resolve()
    if not template.exists():
        raise FileNotFoundError(template)
    output.parent.mkdir(parents=True, exist_ok=True)

    markdown = prepare_markdown()
    with tempfile.TemporaryDirectory(prefix="tfm_docx_") as temp_dir_name:
        temp_dir = Path(temp_dir_name)
        body_path = temp_dir / "body.docx"
        master_path = temp_dir / "master.docx"
        merged_path = temp_dir / "merged.docx"

        pandoc_body(markdown, template, body_path)
        trim_template_and_fill_cover(template, master_path, args)

        master = Document(master_path)
        composer = Composer(master)
        composer.append(Document(body_path))
        composer.save(merged_path)

        document = Document(merged_path)
        configure_styles(document)
        move_table_captions_below_tables(document)
        insert_indices(document)
        configure_headings_and_paragraphs(document)
        format_tables(document)
        format_figures(document)
        enable_field_updates(document)
        clean_properties(document, args.title)
        checks = validate(document)
        document.save(output)

    clean_package_properties(output, args.title)
    # Re-open after the final ZIP-level metadata cleanup.
    validate(Document(output))
    return checks


def main() -> None:
    args = parse_args()
    checks = build(args)
    print(f"DOCX generated: {args.output.expanduser().resolve()}")
    for key, value in checks.items():
        print(f"  {key}: {value}")


if __name__ == "__main__":
    main()
