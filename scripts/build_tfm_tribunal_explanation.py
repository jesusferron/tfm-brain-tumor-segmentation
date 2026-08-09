#!/usr/bin/env python3
"""Generate the detailed slide-by-slide oral explanation for the TFM tribunal."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "slideshow" / "TFM_explicacion_diapositivas_tribunal_Jesus_Ferron_Rubio.docx"

FONT = "Arial"
ORANGE = "E64F12"
PALE = "FDECE8"
INK = RGBColor(0x16, 0x17, 0x1A)
MID = RGBColor(0x61, 0x64, 0x68)


SLIDES = [
    {
        "number": 1,
        "title": "Segmentación de tumores cerebrales en resonancia magnética multimodal",
        "time": "0:30",
        "objective": "Presentar el tema, el alcance y la pregunta que guía el trabajo.",
        "speech": [
            "Buenos días. Soy Jesús Ferrón Rubio y voy a presentar mi Trabajo de Fin de Máster, centrado en la segmentación tridimensional de gliomas sobre resonancia magnética multimodal. El trabajo utiliza cuatro secuencias complementarias —T1n, T1c, T2w y FLAIR— y estudia cómo combinarlas antes de una red de segmentación.",
            "La contribución experimental no consiste en proponer una herramienta clínica terminada, sino en comprobar de forma controlada si una ponderación adaptativa y ligera de esas modalidades mejora la concatenación directa. Para responder con mayor fiabilidad, las configuraciones principales se repiten con tres semillas y se evalúan sobre una partición final común.",
        ],
        "key": "El TFM evalúa una decisión concreta de fusión multimodal dentro de un sistema reproducible de segmentación 3D.",
        "transition": "Antes de entrar en el problema clínico y técnico, voy a resumir cómo se construye la respuesta experimental.",
        "precision": "No presentar el sistema como método de diagnóstico, cribado o producto listo para uso clínico.",
    },
    {
        "number": 2,
        "title": "Cómo se responde a la pregunta del TFM",
        "time": "0:40",
        "objective": "Dar al tribunal el contexto mínimo y el mapa lógico de la defensa.",
        "speech": [
            "La defensa sigue tres pasos. El punto de partida es un problema multimodal: ninguna secuencia de resonancia muestra por sí sola toda la lesión, por lo que hay que combinar T1n, T1c, T2w y FLAIR para delimitar las regiones ET, TC y WT.",
            "Después planteo un experimento controlado. Mantengo fija la Residual U-Net 3D, los datos, las particiones y el entrenamiento, y cambio únicamente la regla de fusión de entrada. Finalmente, la decisión no se basa en una ejecución aislada: comparo tres semillas mediante Dice y HD95 y compruebo si cualquier mejora es medible y consistente. A partir de aquí recorreré el problema, el diseño, los resultados y sus límites.",
        ],
        "key": "La pregunta se responde aislando la fusión como variable y exigiendo consistencia entre repeticiones.",
        "transition": "Con este mapa, el primer paso es entender por qué la segmentación de un glioma es especialmente exigente.",
        "precision": "Todavía no adelantar la conclusión ni interpretar los valores de Dice; esta diapositiva solo orienta el relato.",
    },
    {
        "number": 3,
        "title": "Delimitar un glioma es una tarea tridimensional y multimodal",
        "time": "0:55",
        "objective": "Explicar la dificultad clínica y computacional de la segmentación.",
        "speech": [
            "Segmentar no significa decidir si existe o no un tumor, sino asignar una región a cada vóxel de un volumen tridimensional. La delimitación manual requiere experiencia, consume tiempo y puede variar entre observadores, especialmente cuando los límites son difusos o la lesión es heterogénea.",
            "Desde el punto de vista computacional aparecen cuatro retos: trabajar con volúmenes 3D, combinar información multimodal, representar subregiones con aspecto diferente y generalizar ante variaciones de escáner, protocolo o centro. La automatización puede apoyar una delimitación más rápida y consistente, pero una buena métrica dentro de BraTS no demuestra por sí sola utilidad clínica ni rendimiento en otros hospitales.",
        ],
        "key": "El modelo debe producir una máscara tridimensional precisa a partir de señales complementarias y heterogéneas.",
        "transition": "Esa complementariedad se entiende mejor observando qué aporta cada modalidad y qué regiones debe devolver el sistema.",
        "precision": "Diferenciar claramente segmentación, detección y diagnóstico; este TFM aborda únicamente la primera.",
    },
    {
        "number": 4,
        "title": "Cuatro secuencias de entrada; tres regiones anidadas de salida",
        "time": "1:05",
        "objective": "Definir las entradas multimodales y las regiones objetivo de BraTS.",
        "speech": [
            "El sistema recibe cuatro volúmenes alineados. T1n aporta una referencia anatómica; T1c ayuda a visualizar el tejido que realza después del contraste; T2w es sensible al contenido de agua; y FLAIR suprime el líquido cefalorraquídeo y facilita observar alteraciones periféricas. Estas descripciones indican tendencias, no una correspondencia exclusiva entre modalidad y región.",
            "La salida contiene tres máscaras anidadas. ET representa el tumor realzante, TC el núcleo tumoral y WT el tumor completo. La relación es ET dentro de TC y TC dentro de WT. Por tanto, el reto no consiste solo en introducir cuatro canales, sino en combinarlos de manera que la red aproveche su complementariedad para delimitar correctamente las tres regiones.",
        ],
        "key": "Las modalidades aportan contrastes complementarios y las salidas ET, TC y WT forman regiones anidadas.",
        "transition": "A partir de esta necesidad de combinar modalidades se formula la pregunta central del trabajo.",
        "precision": "Evitar afirmar que una modalidad determina por sí sola una región; la interpretación siempre es conjunta.",
    },
    {
        "number": 5,
        "title": "Pregunta de investigación",
        "time": "1:00",
        "objective": "Formular la hipótesis y delimitar qué significa mejorar.",
        "speech": [
            "La pregunta es si una ponderación explícita, ligera y dependiente de la entrada puede mejorar de forma consistente la segmentación frente a la concatenación directa. La palabra explícita significa que se obtiene un peso observable por modalidad. Dependiente de la entrada significa que esos pesos pueden cambiar según el parche o la ventana procesada.",
            "La comparación se mantiene acotada: se usa la misma Residual U-Net 3D y el mismo presupuesto de entrenamiento. Una mejora no se acepta solo porque una semilla obtenga un valor ligeramente superior; debe aparecer de forma suficientemente estable y sin un aumento desproporcionado del coste. Esta formulación permite que un resultado negativo también responda válidamente a la hipótesis.",
        ],
        "key": "La hipótesis exige una mejora reproducible frente a concatenación, no un máximo aislado.",
        "transition": "Para atribuir cualquier diferencia a la fusión, el diseño experimental congela el resto de componentes.",
        "precision": "No formular la hipótesis como si la mejora estuviera garantizada; el objetivo es comprobarla.",
    },
    {
        "number": 6,
        "title": "Un experimento de ablación: una variable cambia y el resto se congela",
        "time": "1:15",
        "objective": "Justificar la validez interna de la comparación de fusión.",
        "speech": [
            "Esta es la base metodológica de la contribución principal. En las cuatro configuraciones mantengo fija una Residual U-Net 3D de aproximadamente 1,19 millones de parámetros. También se mantienen las mismas particiones, transformaciones, función de pérdida, presupuesto de 15.000 pasos y evaluación sobre los mismos 243 estudios de test.",
            "Solo cambia el bloque situado antes del codificador: concatenación directa, ponderación global, compuerta adaptativa basada en la media y compuerta basada en media más desviación típica. Cada configuración se entrena con tres semillas, por lo que la ablación suma doce ejecuciones. Attention U-Net, Swin-UNETR y nnU-Net se utilizarán después como referencias contextuales, pero no pertenecen a esta comparación causal porque cambian más elementos del sistema.",
        ],
        "key": "Cuatro estrategias por tres semillas permiten estudiar la fusión sin confundirla con un cambio de arquitectura.",
        "transition": "El siguiente elemento que debe permanecer común y trazable es el conjunto de datos.",
        "precision": "Distinguir siempre la ablación controlada de la comparación descriptiva entre arquitecturas.",
    },
    {
        "number": 7,
        "title": "BraTS-GLI 2024: control de calidad y particiones reproducibles",
        "time": "1:10",
        "objective": "Explicar la procedencia, verificación y reparto de los estudios.",
        "speech": [
            "Se utilizaron 1.621 estudios post-tratamiento de BraTS-GLI 2024 con las cuatro modalidades y máscara. El control de calidad verificó para todos ellos la presencia de archivos, la geometría, las transformaciones afines y las etiquetas permitidas. No se detectaron casos problemáticos en ese control.",
            "La partición fue 70 % para entrenamiento, 15 % para validación y 15 % para test: 1.135, 243 y 243 estudios. Se estratificó por origen, presencia de ET y cuartil de volumen de WT. El principal límite es que el reparto se hizo por identificador de estudio y no por paciente. Por ello, el test es válido para comparar las configuraciones sobre un conjunto común, pero no constituye una estimación independiente de generalización a pacientes completamente nuevos.",
        ],
        "key": "El reparto es reproducible y común a los modelos, pero su unidad es el estudio, no el paciente.",
        "transition": "Con los datos verificados y congelados, el pipeline organiza el entrenamiento, la inferencia y la evaluación.",
        "precision": "Exponer el límite del split por estudio de forma directa; ocultarlo debilitaría la defensa metodológica.",
    },
    {
        "number": 8,
        "title": "Un flujo común de datos, predicción y evaluación",
        "time": "1:00",
        "objective": "Mostrar la reproducibilidad y la separación entre las rutas MONAI y nnU-Net.",
        "speech": [
            "El flujo comienza con BraTS-GLI, el control de calidad y las particiones versionadas. La ruta principal usa MONAI para cargar NIfTI, normalizar por modalidad, transformar las etiquetas en ET, TC y WT, aplicar aumento durante entrenamiento y realizar inferencia por ventana deslizante.",
            "Residual U-Net con las cuatro fusiones, Attention U-Net y Swin-UNETR comparten la interfaz MONAI. nnU-Net mantiene su propio preprocesamiento y entrenamiento, pero sus predicciones se llevan al evaluador común. Además, predict y evaluate son operaciones separadas, por lo que las métricas pueden recalcularse sin repetir inferencia. Los YAML, la CLI, los manifiestos y los CSV/JSON finales proporcionan trazabilidad desde una configuración hasta su resultado.",
        ],
        "key": "Las rutas de entrenamiento pueden diferir, pero las predicciones convergen en una evaluación común y trazable.",
        "transition": "Dentro de la rama residual, el componente que se modifica es exactamente el bloque de fusión de entrada.",
        "precision": "No afirmar que nnU-Net comparte el mismo entrenamiento; comparte datos de origen y evaluación final, no su pipeline interno.",
    },
    {
        "number": 9,
        "title": "Cuatro estrategias de fusión antes del mismo codificador",
        "time": "1:30",
        "objective": "Explicar técnicamente las variantes de fusión comparadas.",
        "speech": [
            "La concatenación entrega los cuatro canales directamente a la U-Net. Esto no implica pesos idénticos: la primera convolución ya puede aprender filtros diferentes, aunque no produce un coeficiente separado e interpretable por modalidad. La ponderación global añade cuatro escalares aprendidos y compartidos por todas las entradas.",
            "Las compuertas adaptativas calculan descriptores globales del tensor. La primera usa la media de cada modalidad; la segunda añade la desviación típica. Un pequeño MLP produce cuatro valores, un softmax los normaliza y los pesos reescalan los canales antes del codificador. El incremento es de 76 o 108 parámetros. Es importante remarcar que no se trata de atención espacial: cada peso permanece constante dentro del parche o ventana, aunque puede depender de la entrada procesada.",
        ],
        "key": "Las compuertas son ligeras y dependientes de la entrada, pero no modelan variaciones espaciales dentro del tumor.",
        "transition": "Con las cuatro variantes definidas, el protocolo establece cómo entrenarlas y cómo decidir cuál funciona mejor.",
        "precision": "No describir la compuerta como atención por vóxel o por subregión; su ponderación es global dentro de cada entrada.",
    },
    {
        "number": 10,
        "title": "Protocolo común: repetir, congelar y medir por región",
        "time": "1:05",
        "objective": "Explicar la selección de checkpoints y las métricas finales.",
        "speech": [
            "Cada ejecución se entrena durante 15.000 pasos. La validación se utiliza para monitorizar y seleccionar el mejor checkpoint. Una vez fijada la configuración, el test no interviene en esa selección y se procesa mediante ventana deslizante sobre los 243 estudios.",
            "La métrica principal es Dice, que mide el solapamiento entre la máscara predicha y la referencia: cuanto más próximo a uno, mejor. HD95 mide la discrepancia de superficie en milímetros y se interpreta al contrario: un valor menor es mejor. Ambas se calculan para ET, TC y WT. Para las configuraciones repetidas se informa de la media y la desviación estándar poblacional de las tres semillas, lo que permite observar no solo rendimiento, sino estabilidad.",
        ],
        "key": "El test se reserva para el final y la decisión combina solapamiento, distancia de superficie y variabilidad.",
        "transition": "Aplicando este protocolo aparece el resultado central de la ablación.",
        "precision": "Tres semillas describen variabilidad, pero no permiten una inferencia estadística fuerte sobre toda la población de entrenamientos.",
    },
    {
        "number": 11,
        "title": "Resultado central: la fusión adaptativa no mejora de forma consistente",
        "time": "1:35",
        "objective": "Responder cuantitativamente a la hipótesis principal.",
        "speech": [
            "La concatenación obtiene un Dice medio de 0,706 con una desviación de 0,005. La ponderación global alcanza prácticamente el mismo valor, 0,706 con desviación de 0,006. Por tanto, aprender cuatro pesos globales no ofrece una ventaja sistemática sobre dejar que la primera convolución procese directamente los canales.",
            "Las dos variantes adaptativas quedan alrededor de 0,59 y presentan una desviación cercana a 0,17. Las barras de error muestran que el problema principal no es una reducción pequeña y uniforme, sino una fuerte inestabilidad. Bajo este protocolo, la respuesta a la hipótesis es negativa: las compuertas ligeras estudiadas no mejoran de manera consistente la concatenación. Esta conclusión se limita a estas configuraciones y no descarta mecanismos espaciales o fusiones intermedias.",
        ],
        "key": "Concatenación y ponderación global son equivalentes; las compuertas adaptativas pierden estabilidad.",
        "transition": "Para entender esa dispersión conviene observar las ejecuciones por semilla.",
        "precision": "Decir «no mejoran bajo el protocolo evaluado», no «la fusión adaptativa nunca funciona».",
    },
    {
        "number": 12,
        "title": "La diferencia está en la estabilidad entre semillas",
        "time": "1:15",
        "objective": "Mostrar que una única ejecución habría producido una interpretación engañosa.",
        "speech": [
            "Concatenación y ponderación global permanecen aproximadamente entre 0,697 y 0,713 en las tres semillas. En cambio, cada compuerta adaptativa tiene dos ejecuciones comparables al control y una caída hasta aproximadamente 0,35. Por eso la media agregada desciende y la desviación estándar aumenta.",
            "Este patrón es metodológicamente importante: si solo se hubiera conservado una de las semillas favorables, se podría haber defendido una mejora aparente. La repetición muestra que esa mejora no es robusta. En un diagnóstico exploratorio sobre cuarenta volúmenes preliminares los pesos variaron poco entre estudios, pero esa prueba no corresponde a los checkpoints finales ni mide exactamente los pesos calculados sobre todos los parches y ventanas. Por tanto, describe una señal, pero no demuestra la causa del colapso.",
        ],
        "key": "El análisis multisemilla transforma una posible mejora aislada en un resultado negativo defendible.",
        "transition": "Las métricas se complementan ahora con una inspección visual de un corte axial.",
        "precision": "No atribuir causalidad al diagnóstico de pesos; su alcance es exploratorio.",
    },
    {
        "number": 13,
        "title": "Corte axial: comparación cualitativa de las segmentaciones",
        "time": "1:10",
        "objective": "Hacer visible qué representan ET, TC y WT y cómo difieren las predicciones.",
        "speech": [
            "Esta figura muestra el mismo corte del estudio BraTS-GLI-02273-100. A la izquierda está la referencia manual y después aparecen nnU-Net, Swin-UNETR, Attention U-Net y la Residual U-Net con concatenación. El rojo representa ET, el naranja TC y el amarillo WT.",
            "Todos los modelos recuperan la extensión principal de la lesión, pero existen diferencias locales en los bordes y en la distribución interna de las regiones. nnU-Net se aproxima especialmente bien a la referencia en este ejemplo, mientras que los modelos MONAI muestran pequeñas variaciones. El caso fue seleccionado post hoc y favorece una visualización clara, de modo que no se utiliza como prueba representativa. Su función es traducir las métricas agregadas a un ejemplo anatómico comprensible.",
        ],
        "key": "Un mismo Dice global puede ocultar diferencias locales relevantes en límites y subregiones.",
        "transition": "El segundo corte cualitativo ilustra específicamente el tipo de error de una ejecución adaptativa débil.",
        "precision": "No generalizar el comportamiento de un modelo a partir de este único estudio seleccionado.",
    },
    {
        "number": 14,
        "title": "Corte axial: ejemplo de una ejecución adaptativa de bajo rendimiento",
        "time": "1:00",
        "objective": "Ilustrar visualmente el fallo observado en una de las semillas adaptativas.",
        "speech": [
            "En el mismo corte comparo la referencia manual, la concatenación de la semilla 20260526 y la compuerta adaptativa de la semilla 20260528, que es la ejecución de bajo rendimiento. La concatenación reproduce de forma razonable la lesión principal y su estructura interna.",
            "En la predicción adaptativa destaca una región ET roja mucho más extensa, que sobreestima el tumor realzante y representa con menor precisión la distribución de TC y WT. La imagen ayuda a entender qué tipo de error contribuye a las métricas bajas de esa ejecución. De nuevo, no constituye evidencia independiente ni explica el mecanismo del fallo; la conclusión procede de los 243 estudios y de la comparación de las tres semillas.",
        ],
        "key": "La corrida débil no solo reduce una cifra: produce una sobresegmentación visible de ET.",
        "transition": "Después de la ablación principal, sitúo sus resultados frente a arquitecturas de referencia más amplias.",
        "precision": "Presentar la imagen como ilustración post hoc, no como selección aleatoria ni análisis causal.",
    },
    {
        "number": 15,
        "title": "Contexto arquitectónico: hay margen más allá de la fusión temprana",
        "time": "1:20",
        "objective": "Situar la ablación dentro de referencias arquitectónicas más potentes.",
        "speech": [
            "nnU-Net obtiene el mayor Dice medio observado, 0,829, aunque corresponde a una sola ejecución y utiliza su propio pipeline auto-configurado. Entre los modelos MONAI, Swin-UNETR alcanza 0,752 y Attention U-Net 0,735. Las variantes residuales estables quedan alrededor de 0,706.",
            "Esta ordenación indica que existe margen de mejora más allá de añadir cuatro pesos en la entrada. Sin embargo, no debe interpretarse como una ablación de arquitectura: nnU-Net cambia preprocesamiento, pérdida, optimización y planificación; además, los modelos se ejecutaron en entornos distintos. La comparación es descriptiva. Swin-UNETR sí aporta la referencia Transformer-UNet incluida en el título, pero el diseño no permite atribuir toda su diferencia únicamente al uso de Transformers.",
        ],
        "key": "El rendimiento depende del sistema completo; una fusión ligera no sustituye arquitectura, entrenamiento y auto-configuración.",
        "transition": "Además del rendimiento, el trabajo documenta el coste paramétrico y los elementos de reproducibilidad.",
        "precision": "No afirmar superioridad estadística ni causal de nnU-Net o Swin-UNETR frente a la familia residual.",
    },
    {
        "number": 16,
        "title": "Ligereza paramétrica y trazabilidad experimental",
        "time": "1:00",
        "objective": "Evaluar el coste añadido y resumir la contribución reproducible.",
        "speech": [
            "Los bloques son realmente ligeros en número de parámetros: la ponderación global añade cuatro, la compuerta por media añade 76 y la variante media más desviación añade 108 sobre una red de 1.190.358 parámetros. En MPS, la concatenación necesitó unas cinco horas por ejecución y las variantes ponderadas entre 5,65 y 6,02 horas.",
            "Estos tiempos son duración de pared, no un benchmark aislado; incluyen posibles diferencias de entrada, salida y validación, por lo que no se puede atribuir todo el incremento al bloque. La contribución reproducible incluye splits y manifiesto, configuraciones YAML, una CLI con cinco operaciones y artefactos CSV/JSON por ejecución. Así, las doce corridas de la ablación pueden rastrearse desde su configuración hasta las métricas finales.",
        ],
        "key": "La ligereza paramétrica está demostrada; la eficiencia completa solo está caracterizada parcialmente.",
        "transition": "Con estas evidencias puedo cerrar la respuesta, los límites y las líneas futuras.",
        "precision": "No equiparar pocos parámetros con menor tiempo o memoria sin una medición homogénea.",
    },
    {
        "number": 17,
        "title": "Conclusión, límites y siguiente paso",
        "time": "1:30",
        "objective": "Cerrar la hipótesis con una respuesta acotada y metodológicamente prudente.",
        "speech": [
            "La respuesta es no bajo el protocolo evaluado. La concatenación es la opción más sencilla y una de las más estables. La ponderación global obtiene un resultado equivalente, y las dos compuertas adaptativas presentan una ejecución débil que impide defender una mejora consistente.",
            "La conclusión queda limitada a tres semillas, 15.000 pasos, estas señales globales y una partición por estudio. No existe validación externa ni demostración de utilidad clínica. Como trabajo futuro, la prioridad sería agrupar por paciente antes de repetir el experimento, validar en datos externos y, solo después, explorar mecanismos con información espacial o fusión intermedia y un mayor número de repeticiones. El valor principal del resultado negativo es mostrar que repetir cambió la interpretación que habría surgido de una única semilla favorable.",
        ],
        "key": "El resultado negativo es válido porque la hipótesis se comprobó de forma controlada y multisemilla.",
        "transition": "Con esto finaliza la exposición y queda abierto el turno de preguntas.",
        "precision": "Mantener siempre la fórmula «bajo el protocolo evaluado» y separar comparación interna de validez clínica.",
    },
    {
        "number": 18,
        "title": "Gracias · Preguntas y discusión",
        "time": "0:20",
        "objective": "Cerrar con una síntesis breve y abrir el diálogo con el tribunal.",
        "speech": [
            "Muchas gracias por su atención. Como síntesis final: la ablación reunió doce ejecuciones, la concatenación alcanzó un Dice medio de 0,706 y las compuertas adaptativas estudiadas no ofrecieron una mejora consistente. Quedo a disposición del tribunal para cualquier pregunta sobre los datos, el diseño, las métricas, las arquitecturas o las limitaciones del trabajo.",
        ],
        "key": "La defensa termina con una respuesta clara, cuantificada y limitada al alcance real del estudio.",
        "transition": "Mantener esta diapositiva visible durante las preguntas.",
        "precision": "Responder distinguiendo siempre resultados de la ablación, referencias contextuales e implicaciones clínicas.",
    },
]


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_labelled_paragraph(doc, label: str, text: str, *, italic=False):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    label_run = paragraph.add_run(label + " ")
    label_run.bold = True
    label_run.font.color.rgb = RGBColor(0xE6, 0x4F, 0x12)
    text_run = paragraph.add_run(text)
    text_run.italic = italic
    return paragraph


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    styles = doc.styles
    styles["Normal"].font.name = FONT
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].font.color.rgb = INK
    styles["Normal"].paragraph_format.line_spacing = 1.08
    styles["Heading 1"].font.name = FONT
    styles["Heading 1"].font.size = Pt(17)
    styles["Heading 1"].font.bold = True
    styles["Heading 1"].font.color.rgb = RGBColor(0xE6, 0x4F, 0x12)
    styles["Heading 1"].paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = "TFM · Explicación diapositiva a diapositiva para el tribunal"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = FONT
        run.font.size = Pt(8)
        run.font.color.rgb = MID
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Jesús Ferrón Rubio · Trabajo Fin de Máster")
    for run in footer.runs:
        run.font.name = FONT
        run.font.size = Pt(8)
        run.font.color.rgb = MID

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(70)
    run = title.add_run("Explicación de la presentación\npara el tribunal")
    run.font.name = FONT
    run.font.size = Pt(27)
    run.font.bold = True
    run.font.color.rgb = INK
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(16)
    run = subtitle.add_run("Trabajo Fin de Máster · Jesús Ferrón Rubio")
    run.font.name = FONT
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0xE6, 0x4F, 0x12)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(8)
    meta.add_run("18 diapositivas · Duración orientativa total: 19:20\nCurso 2025–2026")
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(28)
    note_run = note.add_run(
        "Este documento desarrolla el discurso oral sugerido. Conviene utilizarlo para preparar la defensa, no leerlo literalmente ante el tribunal."
    )
    note_run.italic = True
    note_run.font.color.rgb = MID
    intro = doc.add_heading("Cómo utilizar este documento", level=1)
    doc.add_paragraph(
        "Cada sección corresponde a una diapositiva del PowerPoint. La explicación sugerida está redactada en primera persona para facilitar el ensayo. La idea clave resume lo que debería retener el tribunal; la transición enlaza con la siguiente diapositiva; y la precisión señala un límite o una formulación que conviene respetar durante la defensa."
    )

    for slide in SLIDES:
        heading = doc.add_heading(
            f"Diapositiva {slide['number']:02d} · {slide['title']}", level=1
        )
        heading.paragraph_format.space_before = Pt(14)
        heading.paragraph_format.space_after = Pt(4)
        add_labelled_paragraph(
            doc,
            "Tiempo y objetivo:",
            f"{slide['time']} · {slide['objective']}",
        )

        label = doc.add_paragraph()
        label.paragraph_format.space_before = Pt(4)
        label.paragraph_format.space_after = Pt(2)
        label_run = label.add_run("Explicación oral sugerida")
        label_run.bold = True
        label_run.font.color.rgb = RGBColor(0xE6, 0x4F, 0x12)
        for text in slide["speech"]:
            paragraph = doc.add_paragraph(text)
            paragraph.paragraph_format.space_after = Pt(5)

        table = doc.add_table(rows=1, cols=1)
        table.autofit = True
        cell = table.cell(0, 0)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, PALE)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run("Idea que debe quedar: ")
        run.bold = True
        run.font.color.rgb = RGBColor(0xE6, 0x4F, 0x12)
        paragraph.add_run(slide["key"])

        add_labelled_paragraph(doc, "Transición:", slide["transition"], italic=True)
        add_labelled_paragraph(doc, "Precisión ante el tribunal:", slide["precision"])

        separator = doc.add_paragraph()
        separator.paragraph_format.space_after = Pt(2)
        run = separator.add_run("―" * 36)
        run.font.color.rgb = RGBColor(0xCD, 0xCF, 0xCE)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)
    print(f"slides_explained={len(SLIDES)}")


if __name__ == "__main__":
    build_document()
